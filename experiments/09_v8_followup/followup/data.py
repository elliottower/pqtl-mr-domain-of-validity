"""Indices and keys of the follow-up plan (§Study design, §Measured variables, §Indices), computed
on the V8 stage D join (`stage_d.pipeline.load_joined`): the unique genetic unit and its integrity
check, the colocalization class, B and C, midrank quintiles, S_0.5 and assay_cis_signal_strength."""
import math
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from followup.common import FollowupError
from stage_d.evidence import evidence_state

UNIT_KEY = ("instrument_source", "instrument_assay_id", "gene_ensembl", "outcome_accession")
UNIT_SHARED = ("coloc_run", "not_run_reason", "n_shared", "lead_variant", "genetic_direction")
PP = ("pp_h0", "pp_h1", "pp_h2", "pp_h3", "pp_h4")
PP_TOL = 1e-12
CLASS_THRESHOLD = 0.8
C_MIN_B = 0.01
STRENGTH_CAP = 300.0
ENSG_RE = r"ENSG\d{11}"


def unit_id(df: pd.DataFrame) -> pd.Series:
    return df[list(UNIT_KEY)].astype(str).agg("|".join, axis=1)


def check_unit_integrity(df: pd.DataFrame) -> None:
    """Every row of one unique genetic unit must carry the same colocalization result."""
    uid = unit_id(df)
    for col in UNIT_SHARED:
        n = df[col].astype(object).where(df[col].notna(), "<NA>").groupby(uid).nunique()
        if (n > 1).any():
            raise FollowupError(f"unit {n.index[n > 1][0]} has {n.max()} different values of {col}")
    for col in PP:
        v = df[col].astype(float)
        spread = (v.groupby(uid).max() - v.groupby(uid).min()).fillna(0.0)
        mixed = v.isna().groupby(uid).nunique() > 1
        if (spread > PP_TOL).any() or mixed.any():
            raise FollowupError(f"unit {spread.index[(spread > PP_TOL) | mixed][0]} has different {col} values")


def check_ensembl(ids: pd.Series, what: str) -> None:
    bad = ~ids.astype(str).str.fullmatch(ENSG_RE)
    if bad.any():
        raise FollowupError(f"{what}: {int(bad.sum())} ids are not ENSG + 11 digits (first {ids[bad].iloc[0]})")


def coloc_class(df: pd.DataFrame) -> pd.Series:
    """§Indices: not_run, then the first of shared, distinct, outcome_only, no_outcome_signal that
    holds, else unresolved_<largest posterior> (H0+H1, H2, H3, H4; ties to the earlier)."""
    ran = df["coloc_run"].astype(bool)
    pp = df.loc[ran, list(PP)].astype(float)
    if pp.isna().any().any():
        raise FollowupError(f"{int(pp.isna().any(axis=1).sum())} completed runs have a missing posterior")
    out = pd.Series("not_run", index=df.index, dtype=object)
    groups = pd.DataFrame({"h0h1": pp["pp_h0"] + pp["pp_h1"], "h2": pp["pp_h2"], "h3": pp["pp_h3"], "h4": pp["pp_h4"]})
    largest = groups.idxmax(axis=1)          # idxmax returns the first maximum: the earlier label wins ties
    cls = "unresolved_" + largest
    cls = cls.mask(groups["h0h1"] >= CLASS_THRESHOLD, "no_outcome_signal")
    cls = cls.mask(pp["pp_h2"] >= CLASS_THRESHOLD, "outcome_only")
    cls = cls.mask(pp["pp_h3"] >= CLASS_THRESHOLD, "distinct")
    cls = cls.mask(pp["pp_h4"] >= CLASS_THRESHOLD, "shared")
    out.loc[ran] = cls
    return out


def b_and_c(df: pd.DataFrame) -> tuple[pd.Series, pd.Series]:
    """B = PP.H3 + PP.H4 on completed runs; C = PP.H4 / B where B >= 0.01 (C = 1 when PP.H3 = 0 < PP.H4)."""
    ran = df["coloc_run"].astype(bool)
    h3, h4 = df["pp_h3"].astype(float), df["pp_h4"].astype(float)
    b = (h3 + h4).where(ran)
    c = (h4 / b).where(ran & (b >= C_MIN_B))
    return b, c


def midrank_quintile(x: pd.Series) -> pd.Series:
    """Quintile 1-5 of the non-missing values: midrank / n cut at 0.2, 0.4, 0.6, 0.8, a value exactly
    at a cut in the lower quintile; missing stays missing."""
    v = x.astype(float)
    ok = v.notna()
    r = v[ok].rank(method="average") / ok.sum()
    q = pd.Series(pd.NA, index=x.index, dtype="Int64")
    q[ok] = np.searchsorted([0.2, 0.4, 0.6, 0.8], r.to_numpy(), side="left") + 1
    return q


def s05_state(df: pd.DataFrame) -> np.ndarray:
    """S_0.5: the V8 state rule at PP.H4 >= 0.5, directions unchanged (a run not completed has no PP.H4)."""
    return evidence_state(df["pp_h4"].where(df["coloc_run"].astype(bool)), df["genetic_direction"], df["direction"],
                          threshold=0.5)


# ---- assay_cis_signal_strength ------------------------------------------------------------------

def _norm(name) -> str:
    return " ".join(str(name).split())


def _sheet_rows(path: Path, sheet: str, header_row: int, columns: dict[str, str]) -> list[dict]:
    wb = load_workbook(path, read_only=True)
    if sheet not in wb.sheetnames:
        raise FollowupError(f"{path.name}: no sheet {sheet}")
    rows = wb[sheet].iter_rows(min_row=header_row, values_only=True)
    header = [_norm(h) for h in next(rows)]
    idx = {}
    for key, name in columns.items():
        if _norm(name) not in header:
            raise FollowupError(f"{path.name} {sheet}: no column {name!r}")
        idx[key] = header.index(_norm(name))
    return [{k: r[i] for k, i in idx.items()} for r in rows if r is not None and any(c is not None for c in r)]


def _mlog10(value) -> tuple[float | None, bool]:
    """(−log10 p, truncated): infinite or > 300 is set to 300 and flagged; unreadable is None."""
    try:
        v = float(value)
    except (TypeError, ValueError):
        if str(value).strip().lower() in ("inf", "infinity"):
            return STRENGTH_CAP, True
        return None, False
    if math.isnan(v):
        return None, False
    if math.isinf(v) or v > STRENGTH_CAP:
        return STRENGTH_CAP, True
    return v, False


def max_by_assay(records: list[tuple[str, object]]) -> dict[str, dict]:
    """assay -> {value: max over the assay's rows, rows, truncated}."""
    out: dict[str, dict] = {}
    for assay, raw in records:
        v, trunc = _mlog10(raw)
        if v is None:
            continue
        cur = out.get(assay)
        if cur is None:
            out[assay] = {"value": v, "rows": 1, "truncated": trunc}
            continue
        cur["rows"] += 1
        if v > cur["value"]:
            cur["value"], cur["truncated"] = v, trunc
        elif v == cur["value"]:
            cur["truncated"] = cur["truncated"] or trunc
    return out


SUN_ST9 = ("ST9", 5, {"protein": "UKBPPP ProteinID", "mlog10p": "log10(p) (discovery)", "cis": "cis/trans"})
FERK_ST08 = ("ST08", 3, {"seqid": "SeqId", "mlog10p": "-Log10(P) (unadj.)"})


def ukbppp_strength(path: Path) -> dict[str, dict]:
    rows = _sheet_rows(path, *SUN_ST9)
    recs = []
    for r in rows:
        if str(r["cis"]).strip().lower() != "cis":
            continue
        oid = next((f for f in str(r["protein"]).split(":") if f.startswith("OID")), None)
        if oid is None:
            raise FollowupError(f"UKBPPP ProteinID {r['protein']!r} carries no OID field")
        recs.append((oid, r["mlog10p"]))
    return max_by_assay(recs)


def decode_strength(path: Path) -> dict[str, dict]:
    return max_by_assay([(str(r["seqid"]).strip(), r["mlog10p"]) for r in _sheet_rows(path, *FERK_ST08)])


def assay_strength(df: pd.DataFrame, by_source: dict[str, dict[str, dict]]) -> tuple[pd.Series, dict]:
    """Per hypothesis, from its exact assay in its own source's table; never another source."""
    values = pd.Series(np.nan, index=df.index)
    report = {}
    for source, g in df.groupby("instrument_source"):
        table = by_source.get(source)
        assays = g["instrument_assay_id"].astype(str)
        hits = assays.map(lambda a: table.get(a) if table else None)
        values[g.index] = hits.map(lambda h: h["value"] if h else np.nan)
        uniq = assays.unique()
        report[source] = {
            "table": "none" if table is None else "pinned",
            "assays": len(uniq), "assays_matched": int(sum(1 for a in uniq if table and a in table)),
            "assays_unmatched": int(sum(1 for a in uniq if not (table and a in table))),
            "assays_multi_row": int(sum(1 for a in uniq if table and a in table and table[a]["rows"] > 1)),
            "assays_truncated_at_300": int(sum(1 for a in uniq if table and a in table and table[a]["truncated"])),
            "hypotheses_missing": int(hits.isna().sum()),
        }
    return values, report
