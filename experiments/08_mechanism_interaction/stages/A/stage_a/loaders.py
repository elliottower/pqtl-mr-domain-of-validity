"""File readers for stage A. Each reads only the columns or sheets the rules need.

Blinding (PREREG §Explanation of foreknowledge, "Phase blinding in stage A"; §Additional
blinding): `clinical_indication.maxClinicalStage` is reduced to a boolean Phase II+ flag on
read and the phase string is dropped before the table leaves this module; `clinicalReportIds`
is never read. `drug_molecule` is read without `maximumClinicalStage`. The pilot
classification files are read with gene and disease only; the EpiGraphDB catalogue with
gene_symbol only; the Karim workbook with sheet ST17 (launched pairs), gene and MeSH ID only.
No Eldjarn 2023 sheet is read by stage A.
"""
import io
import json
import zipfile
from pathlib import Path

import openpyxl
import pandas as pd
import pyarrow.parquet as pq

PHASE2_PLUS = frozenset({"PHASE_2", "PHASE_2_3", "PHASE_3", "PREAPPROVAL", "APPROVAL"})
KNOWN_STAGES = frozenset({"PRECLINICAL", "IND", "EARLY_PHASE_1", "PHASE_1", "PHASE_1_2", "PHASE_2",
                          "PHASE_2_3", "PHASE_3", "PREAPPROVAL", "APPROVAL", "UNKNOWN"})
DISEASE_COLUMNS = ["id", "name", "dbXRefs", "exactSynonyms", "descendants", "ancestors", "obsoleteTerms",
                   "therapeuticAreas", "ontology"]
FINNGEN_COLUMNS = ["studyId", "projectId", "studyType", "traitFromSource", "traitFromSourceMappedIds",
                   "diseaseIds", "nCases", "nControls", "cohorts", "hasSumstats"]
HPA_BLOOD = "Secreted to blood"
KARIM_LAUNCHED_SHEET = "ST17 - pqtl_success_ti_pairs"
SUN2018_SHEET = "ST4 - pQTL summary"


class UnknownClinicalStage(ValueError):
    """clinical_indication carries a stage string outside the registered vocabulary."""


class MissingColumn(ValueError):
    """An input table lacks a column a rule reads."""


def _require(columns, needed, what: str) -> None:
    missing = [c for c in needed if c not in columns]
    if missing:
        raise MissingColumn(f"{what}: missing {missing}")


def load_clinical_indication(path: Path) -> pd.DataFrame:
    """drugId, diseaseId, phase2plus. The phase string never leaves this function."""
    t = pq.read_table(path, columns=["drugId", "diseaseId", "maxClinicalStage"]).to_pandas()
    unknown = sorted(set(t["maxClinicalStage"].dropna()) - KNOWN_STAGES)
    if unknown:
        raise UnknownClinicalStage(f"unregistered maxClinicalStage values: {unknown}")
    out = pd.DataFrame({"drugId": t["drugId"], "diseaseId": t["diseaseId"],
                        "phase2plus": t["maxClinicalStage"].isin(PHASE2_PLUS)})
    del t
    return out.drop_duplicates().reset_index(drop=True)


def load_mechanisms(path: Path) -> pd.DataFrame:
    moa = pq.read_table(path, columns=["mechanismOfAction", "actionType", "chemblIds", "targets",
                                       "targetType"]).to_pandas()
    moa = moa.explode("chemblIds").explode("targets").dropna(subset=["chemblIds", "targets"])
    return moa.rename(columns={"chemblIds": "moa_drug", "targets": "ensembl"}).reset_index(drop=True)


def load_molecules(path: Path) -> pd.DataFrame:
    return pq.read_table(path, columns=["id", "drugType", "parentId"]).to_pandas()


def load_disease(path: Path) -> pd.DataFrame:
    return pq.read_table(path, columns=DISEASE_COLUMNS).to_pandas()


def load_finngen_study(path: Path) -> pd.DataFrame:
    cols = [c for c in FINNGEN_COLUMNS if c in pq.read_schema(path).names]
    _require(cols, ["studyId", "nCases", "nControls"], "OT study index")
    return pq.read_table(path, columns=cols).to_pandas()


def load_hgnc(path: Path) -> tuple[dict[str, str], dict[str, set[str]]]:
    """(Ensembl -> approved symbol, UniProt accession -> approved symbols)."""
    hg = pd.read_csv(path, sep="\t", dtype=str, usecols=["symbol", "status", "ensembl_gene_id", "uniprot_ids"])
    hg = hg[hg["status"] == "Approved"]
    ens2sym = dict(zip(hg["ensembl_gene_id"].dropna(), hg.loc[hg["ensembl_gene_id"].notna(), "symbol"]))
    uni2sym: dict[str, set[str]] = {}
    for s, u in zip(hg["symbol"], hg["uniprot_ids"].fillna("")):
        for acc in u.split("|"):
            if acc:
                uni2sym.setdefault(acc, set()).add(s)
    return ens2sym, uni2sym


def load_hpa_blood(path: Path) -> set[str]:
    """Ensembl IDs whose HPA secretome location is "Secreted to blood"."""
    with zipfile.ZipFile(path) as z:
        hpa = pd.read_csv(io.BytesIO(z.read("proteinatlas.tsv")), sep="\t", dtype=str,
                          usecols=["Ensembl", "Secretome location"])
    return set(hpa.loc[hpa["Secretome location"].fillna("") == HPA_BLOOD, "Ensembl"].dropna())


def load_uniprot_secreted(paths: list[Path]) -> set[str]:
    """Primary gene names of reviewed human entries with KW-0964 or SL-0243."""
    out: set[str] = set()
    for p in paths:
        t = pd.read_csv(p, sep="\t", dtype=str)
        _require(t.columns, ["Gene Names (primary)"], str(p))
        out |= set(t["Gene Names (primary)"].dropna())
    return out


def load_ukbppp_cis(path: Path) -> list[str]:
    entries = json.loads(Path(path).read_text())
    if not isinstance(entries, list):
        raise ValueError(f"{path}: expected a JSON list of assay names")
    return [str(e) for e in entries]


def load_olink_map(path: Path) -> pd.DataFrame:
    return pd.read_csv(path, sep="\t", dtype=str,
                       usecols=["Assay", "OlinkID", "UniProt", "UniProt2", "HGNC.symbol", "ensembl_id"])


def _sheet(path: Path, sheet: str, header_row: int, columns: dict[str, str],
           reset_dimensions: bool = False) -> pd.DataFrame:
    """Rows of one sheet below `header_row` (1-based), keeping `columns` (header -> name), each
    cell as openpyxl gives it (dtype object: an empty cell stays None, never NaN).

    The default reads the sheet as count_v4.py and count_all_indications.py do: read-only,
    `iter_rows(min_row=header_row, values_only=True)` over the dimensions the workbook records,
    without `reset_dimensions`, each cell taken by its header's position. `reset_dimensions=True`
    re-scans the sheet instead and tolerates short rows; only the Karim workbook, which neither
    script reads, uses it."""
    wb = openpyxl.load_workbook(path, read_only=True)
    try:
        ws = wb[sheet]
        if reset_dimensions:
            ws.reset_dimensions()
        rows = ws.iter_rows(min_row=header_row, values_only=True)
        header = [str(c) if c is not None else "" for c in next(rows)]
        _require(header, list(columns), f"{path}:{sheet}")
        idx = {name: header.index(h) for h, name in columns.items()}
        if reset_dimensions:
            data = [{name: (r[i] if i < len(r) else None) for name, i in idx.items()} for r in rows]
        else:
            data = [{name: r[i] for name, i in idx.items()} for r in rows]
    finally:
        wb.close()
    return pd.DataFrame(data, columns=list(columns.values()), dtype=object)


def load_decode(path: Path) -> tuple[pd.DataFrame, pd.DataFrame]:
    """Ferkingstad 2021 ST01 (aptamer ids) and ST02 (gene, SeqId, cis/trans), read as count_v4.py
    lines 227-240 read them (header on row 3)."""
    st01 = _sheet(path, "ST01", 3, {"SeqId": "SeqId", "Gene": "Gene", "UniProt": "UniProt",
                                    "Ensembl.Gene.ID": "Ensembl", "Included in\nanalysis": "Included"})
    st02 = _sheet(path, "ST02", 3, {"gene\n (prot.)": "gene", "SeqId": "SeqId", "cis/\ntrans": "cis_trans"})
    return st01, st02


def load_sun2018_st4(path: Path) -> pd.DataFrame:
    """Sun 2018 ST4, read as count_v4.py lines 270-273 read it (header on row 5)."""
    return _sheet(path, SUN2018_SHEET, 5, {"SOMAmer ID": "somamer_id", "UniProt": "UniProt",
                                           "cis/ trans": "cis_trans"})


def load_epigraphdb_genes(path: Path) -> set[str]:
    epi = pd.read_csv(path, usecols=["gene_symbol"])
    # count_v4.py line 293: every ';'-separated token, an empty one included
    return {g for s in epi["gene_symbol"].dropna().astype(str) for g in s.split(";")}


def load_gwas_catalog(paths: dict[str, Path]) -> list[tuple[pd.DataFrame, pd.DataFrame, str]]:
    """`paths`: studies, ancestries, unpublished_studies, unpublished_ancestries."""
    frames = []
    for st_key, an_key, kind in (("studies", "ancestries", "published"),
                                 ("unpublished_studies", "unpublished_ancestries", "unpublished")):
        st = pd.read_csv(paths[st_key], sep="\t", dtype=str, quoting=3, on_bad_lines="skip")
        an = pd.read_csv(paths[an_key], sep="\t", dtype=str, quoting=3, on_bad_lines="skip")
        st.columns = [c.strip().upper() for c in st.columns]
        an.columns = [c.strip().upper() for c in an.columns]
        frames.append((st, an, kind))
    return frames


def load_opengwas(path: Path) -> dict:
    info = json.loads(Path(path).read_text())
    if not isinstance(info, dict):
        raise ValueError(f"{path}: expected the gwasinfo listing as a JSON object")
    return info


def load_pilot_tables(frozen_candidates: Path, classifications: list[Path]) -> tuple[pd.DataFrame, list[pd.DataFrame]]:
    frozen = pd.read_csv(frozen_candidates, usecols=["disease", "ot_disease_id"]).drop_duplicates()
    return frozen, [pd.read_csv(p, usecols=["gene", "disease"]) for p in classifications]


def load_karim_launched(path: Path) -> pd.DataFrame:
    return _sheet(path, KARIM_LAUNCHED_SHEET, 3, {"gene": "gene", "indication_mesh_id": "indication_mesh_id"},
                  reset_dimensions=True)
