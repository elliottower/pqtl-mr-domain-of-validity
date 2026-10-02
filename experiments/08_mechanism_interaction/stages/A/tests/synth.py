"""A small synthetic universe for stage A tests. Every ID, name and number is invented.

The world is built once as raw tables shaped like the real inputs. `world_inputs` turns it into
`StageAInputs` through the same pure functions the CLI uses; `write_world` serializes it to the
real file formats so `run_stage_a.main` can be run end to end.

Planted cases (gene, indication, drug):
  GA  MONDO_1001 AB1   antibody inhibitor, blood-secreted target -> aligned, S1
  GA  MONDO_3000 AB1   Karim launched (MeSH D000001), oncology, S1
  GA  MONDO_5000 AB1   pilot key -> heldout false, written for the pooled set
  GB  MONDO_1001 SM2   small-molecule inhibitor -> blocking, S1, deCODE instrument
  GB  MONDO_1002 SM2S  salt of SM2 -> program SM2, S1
  GB  MONDO_1000 SM2   ancestor of 1001 and 1002; every program on retained descendants -> dropped
  GB  MONDO_2000 SM3   conflicting rows (inhibitor / disrupting agent) -> other/ambiguous; S21 restores
  GB  MONDO_1001 AB11  antibody inhibitor of a non-secreted target -> other/decrease, S1
  GB  MONDO_1001 SM5   agonist, Phase 1 only -> not eligible
  GB  EFO_4000 / HP_0000001 SM8  indication rule failures
  GC  MONDO_1001 SM7   gene only on a multi-gene UKB-PPP assay -> S20 only
  GE  MONDO_1001 SM4   gene only on inclusive INTERVAL / EpiGraphDB -> S20 only
  GD  MONDO_6000 AB6   neutralizing BINDING AGENT; only a tier-2 outcome GWAS -> S10
  GD  MONDO_7000 AB6   only a tier-1 overlap-unknown outcome GWAS -> S4
  GF  MONDO_8000 AB9   deCODE-only gene; tier-1 GWAS is Icelandic -> tier-2 FinnGen selected, S10
  GG  MONDO_2100 SM10  UKB-PPP assay whose UniProt column holds a label (UniProt2 used); psychiatric only
"""
import io
import json
import zipfile
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path

import openpyxl
import pandas as pd
import pyarrow as pa
import pyarrow.parquet as pq

from stage_a.build import StageAInputs
from stage_a.flags import karim_launched_keys, pilot_keys, pilot_label_map
from stage_a.instruments import InstrumentLists, decode_lists, decode_strict_diagnostic, interval_lists, ukbppp_lists
from stage_a.loaders import DISEASE_COLUMNS, PHASE2_PLUS
from stage_a.outcome_gwas import combine_candidates, finngen_candidates, gwas_catalog_candidates, \
    label_index, opengwas_candidates

DATES = {"ukbppp": date(2030, 3, 3), "decode": date(2020, 2, 2), "interval": date(2010, 1, 1)}

# The quantities of stage_a.registered for the default World, counted by hand from the planted cases.
# Held-out S1: GA/1001 aligned, GA/3000 aligned, GB/1001 blocking, GB/1002 blocking, GG/2100 blocking,
# GB/2000 other, GB/1001 other. GA is the only blood-secreted target among them. MONDO_2000 and
# MONDO_2100 are the neurologic or psychiatric indications. S4 adds GD/7000.
WORLD_COUNTS = {
    "s1_hypotheses": 7, "s1_gene_indication_pairs": 6, "s1_genes": 3, "s1_indications": 5,
    "h1_hypotheses": 5, "h1_aligned_hypotheses": 2, "h1_aligned_genes": 1, "h1_blocking_hypotheses": 3,
    "h1_blocking_genes": 2, "h1_genes_in_both_classes": 0, "s1_other_hypotheses": 2,
    "h4_neuro_psych_hypotheses": 2, "h4_neuro_psych_genes": 2, "h4_other_indication_hypotheses": 5,
    "s4_hypotheses": 8,
    "h1_blood_secreted_aligned_hypotheses": 2, "h1_blood_secreted_aligned_genes": 1,
    "h1_blood_secreted_blocking_hypotheses": 0, "h1_blood_secreted_blocking_genes": 0,
}

GENES = {  # symbol: (ensembl, uniprot)
    "GA": ("ENSG00000000001", "P00001"), "GB": ("ENSG00000000002", "P00002"),
    "GC": ("ENSG00000000003", "P00003"), "GD": ("ENSG00000000004", "P00004"),
    "GE": ("ENSG00000000005", "P00005"), "GF": ("ENSG00000000006", "P00006"),
    "GG": ("ENSG00000000007", "P00007"), "GX": ("ENSG00000000009", "P00009"),
}


def ens(sym: str) -> str:
    return GENES[sym][0]


def disease_row(did: str, name: str, tas=(), desc=(), anc=(), is_ta=False, xrefs=(), syn=(), obs=()) -> dict:
    return {"id": did, "name": name, "dbXRefs": list(xrefs), "exactSynonyms": list(syn),
            "descendants": list(desc), "ancestors": list(anc), "obsoleteTerms": list(obs),
            "therapeuticAreas": list(tas), "ontology": {"isTherapeuticArea": is_ta, "leaf": not desc}}


CV = "EFO_0000319"


def default_diseases() -> list[dict]:
    return [
        disease_row(CV, "cardiovascular disease", tas=[CV], is_ta=True),
        disease_row("MONDO_0045024", "cancer or benign tumor", tas=["MONDO_0045024"], is_ta=True),
        disease_row("MONDO_0005071", "nervous system disorder", tas=["MONDO_0005071"], is_ta=True),
        disease_row("MONDO_0002025", "psychiatric disorder", tas=["MONDO_0002025"], is_ta=True),
        disease_row("EFO_0001444", "measurement", tas=["EFO_0001444"], is_ta=True),
        disease_row("MONDO_1000", "vascular disease", tas=[CV], desc=["MONDO_1001", "MONDO_1002"]),
        disease_row("MONDO_1001", "vascular disease type 1", tas=[CV], anc=["MONDO_1000"]),
        disease_row("MONDO_1002", "vascular disease type 2", tas=[CV], anc=["MONDO_1000"]),
        disease_row("MONDO_2000", "neuropsychiatric disease", tas=["MONDO_0005071", "MONDO_0002025"]),
        disease_row("MONDO_2100", "psychiatric disease", tas=["MONDO_0002025"]),
        disease_row("MONDO_3000", "tumour x", tas=["MONDO_0045024"], xrefs=["MESH:D000001", "UMLS:C1"]),
        disease_row("EFO_4000", "a measurement", tas=["EFO_0001444"]),
        disease_row("HP_0000001", "a phenotype", tas=[CV]),
        disease_row("MONDO_5000", "pilot disease", tas=[CV]),
        disease_row("MONDO_6000", "broad disease", tas=[CV], desc=["MONDO_6001"]),
        disease_row("MONDO_6001", "narrow disease", tas=[CV], anc=["MONDO_6000"]),
        disease_row("MONDO_7000", "unknown overlap disease", tas=[CV]),
        disease_row("MONDO_8000", "iceland disease", tas=[CV], desc=["MONDO_8001"]),
        disease_row("MONDO_8001", "iceland disease subtype", tas=[CV], anc=["MONDO_8000"]),
    ]


def moa(drug: str, action: str, text: str, gene: str, target_type: str = "single protein") -> dict:
    return {"mechanismOfAction": text, "actionType": action, "chemblIds": [drug], "targets": [ens(gene)],
            "targetType": target_type}


def default_mechanisms() -> list[dict]:
    return [
        moa("AB1", "INHIBITOR", "GA inhibitor", "GA"),
        moa("SM2", "INHIBITOR", "GB inhibitor", "GB"),
        moa("SM2", "INHIBITOR", "GB inhibitor", "GB"),  # exact duplicate row
        moa("SM3", "INHIBITOR", "GB inhibitor", "GB"),
        moa("SM3", "DISRUPTING AGENT", "GB disrupting agent", "GB"),
        moa("SM4", "INHIBITOR", "GE inhibitor", "GE"),
        moa("SM5", "AGONIST", "GB agonist", "GB"),
        moa("AB6", "BINDING AGENT", "GD neutralising antibody", "GD"),
        moa("SM7", "INHIBITOR", "GC inhibitor", "GC", "protein complex"),
        moa("SM8", "INHIBITOR", "GB inhibitor", "GB"),
        moa("AB9", "INHIBITOR", "GF inhibitor", "GF"),
        moa("SM10", "ANTAGONIST", "GG antagonist", "GG"),
        moa("AB11", "INHIBITOR", "GB inhibitor", "GB"),
    ]


def default_molecules() -> list[dict]:
    types = {"AB1": "Antibody", "SM2": "Small molecule", "SM2S": "Small molecule", "SM3": "Small molecule",
             "SM4": "Small molecule", "SM5": "Small molecule", "AB6": "Antibody", "SM7": "Small molecule",
             "SM8": "Small molecule", "AB9": "Antibody", "SM10": "Small molecule", "AB11": "Antibody"}
    return [{"id": d, "drugType": t, "parentId": "SM2" if d == "SM2S" else None} for d, t in types.items()]


def default_clinical() -> list[tuple[str, str, str]]:
    return [
        ("AB1", "MONDO_1001", "PHASE_3"), ("AB1", "MONDO_3000", "PHASE_2"), ("AB1", "MONDO_5000", "PHASE_2"),
        ("SM2", "MONDO_1001", "PHASE_2"), ("SM2S", "MONDO_1002", "PHASE_2"), ("SM2", "MONDO_1000", "APPROVAL"),
        ("SM3", "MONDO_2000", "PHASE_2"), ("SM4", "MONDO_1001", "PHASE_2"), ("SM5", "MONDO_1001", "PHASE_1"),
        ("AB6", "MONDO_6000", "PHASE_2"), ("AB6", "MONDO_7000", "PHASE_2_3"), ("SM7", "MONDO_1001", "PHASE_2"),
        ("SM8", "EFO_4000", "PHASE_2"), ("SM8", "HP_0000001", "PHASE_2"), ("AB9", "MONDO_8000", "PREAPPROVAL"),
        ("SM10", "MONDO_2100", "PHASE_2"), ("AB11", "MONDO_1001", "PHASE_2"),
    ]


def gcat_study(acc: str, mondo: str, cohort: str, ncase: int, ncontrol: int, country: str = "Estonia",
               ancestry: str = "European") -> tuple[dict, dict]:
    st = {"STUDY ACCESSION": acc, "DISEASE/TRAIT": f"trait {acc}",
          "MAPPED_TRAIT_URI": f"http://purl.obolibrary.org/obo/{mondo}",
          "INITIAL SAMPLE SIZE": f"{ncase} European ancestry cases, {ncontrol} European ancestry controls",
          "FULL SUMMARY STATISTICS": "yes", "COHORT": cohort}
    an = {"STUDY ACCESSION": acc, "STAGE": "initial", "BROAD ANCESTRAL CATEGORY": ancestry,
          "COUNTRY OF RECRUITMENT": country, "NUMBER OF CASES": str(ncase), "NUMBER OF CONTROLS": str(ncontrol)}
    return st, an


def default_gwas_catalog() -> list[tuple[dict, dict]]:
    return [
        gcat_study("GCST1", "MONDO_1001", "FinnGen|EstBB", 5000, 50000, "Finland"),
        gcat_study("GCST0", "MONDO_1000", "EstBB", 4000, 40000),
        gcat_study("GCST2", "MONDO_2000", "EstBB", 4000, 40000),
        gcat_study("GCST3", "MONDO_3000", "UK Biobank", 100000, 100000, "U.K."),
        gcat_study("GCST5", "MONDO_5000", "EstBB", 4000, 40000),
        gcat_study("GCST61", "MONDO_6001", "EstBB", 4000, 40000),
        gcat_study("GCST7", "MONDO_7000", "", 4000, 40000),
        gcat_study("GCST8", "MONDO_8000", "", 90000, 90000, "Iceland"),
        gcat_study("GCST21", "MONDO_2100", "EstBB", 4000, 40000),
        gcat_study("GCSTSMALL", "MONDO_1001", "EstBB", 500, 500),
        gcat_study("GCSTEAS", "MONDO_1001", "EstBB", 90000, 90000, "Japan", "East Asian"),
    ]


def default_finngen() -> list[dict]:
    def fg(sid: str, mondo: str, ncase: int, ncontrol: int) -> dict:
        return {"studyId": sid, "projectId": "FINNGEN_R12", "studyType": "gwas", "traitFromSource": f"trait {sid}",
                "traitFromSourceMappedIds": [mondo], "diseaseIds": [mondo], "nCases": ncase, "nControls": ncontrol,
                "cohorts": ["FinnGen"], "hasSumstats": True}
    return [fg("FINNGEN_R12_C3", "MONDO_3000", 3000, 300000), fg("FINNGEN_R12_X2", "MONDO_1002", 3000, 300000),
            fg("FINNGEN_R12_Y81", "MONDO_8001", 3000, 300000),
            {"studyId": "GCST_OT_ONLY", "projectId": "GCST", "studyType": "gwas", "traitFromSource": "x",
             "traitFromSourceMappedIds": ["MONDO_1001"], "diseaseIds": ["MONDO_1001"], "nCases": 9, "nControls": 9,
             "cohorts": [], "hasSumstats": True}]


def default_opengwas() -> dict:
    return {
        "ieu-a-1": {"trait": "Vascular disease type 1", "ncase": 9000, "ncontrol": 90000, "population": "European",
                    "consortium": "EstBB", "author": "A", "note": ""},
        "finn-b-X2": {"trait": "finngen x2", "ncase": 3000, "ncontrol": 300000, "population": "European",
                      "consortium": "FinnGen", "author": "B", "note": ""},
        "prot-a-1": {"trait": "protein", "ncase": None, "ncontrol": None, "population": "European",
                     "consortium": "", "author": "", "note": ""},
    }


@dataclass
class World:
    diseases: list[dict] = field(default_factory=default_diseases)
    mechanisms: list[dict] = field(default_factory=default_mechanisms)
    molecules: list[dict] = field(default_factory=default_molecules)
    clinical: list[tuple[str, str, str]] = field(default_factory=default_clinical)
    gwas_catalog: list[tuple[dict, dict]] = field(default_factory=default_gwas_catalog)
    finngen: list[dict] = field(default_factory=default_finngen)
    opengwas: dict = field(default_factory=default_opengwas)
    hpa_blood: tuple[str, ...] = ("GA", "GD")
    uniprot_secreted: tuple[str, ...] = ("GA",)
    ukbppp_cis: tuple[str, ...] = ("GA", "GD", "GC_GX", "GG")
    olink: list[dict] = field(default_factory=lambda: [
        {"Assay": "GA", "OlinkID": "OID001", "UniProt": "P00001", "UniProt2": "P00001", "HGNC.symbol": "GA",
         "ensembl_id": ens("GA")},
        {"Assay": "GD", "OlinkID": "OID004", "UniProt": "P00004", "UniProt2": "P00004", "HGNC.symbol": "GD",
         "ensembl_id": ens("GD")},
        {"Assay": "GC_GX", "OlinkID": "OID003", "UniProt": "P00003_P00009", "UniProt2": "P00003", "HGNC.symbol": "GC",
         "ensembl_id": ens("GC")},
        {"Assay": "GC_GX", "OlinkID": "OID003", "UniProt": "P00003_P00009", "UniProt2": "P00009", "HGNC.symbol": "GX",
         "ensembl_id": ens("GX")},
        {"Assay": "GG", "OlinkID": "OID007", "UniProt": "NTproGG", "UniProt2": "P00007-2", "HGNC.symbol": "GG",
         "ensembl_id": ens("GG")},
    ])
    decode_st01: list[dict] = field(default_factory=lambda: [
        {"SeqId": "1000-1", "Gene": "GB", "UniProt": "P00002", "Ensembl": ens("GB"), "Included": "Yes"},
        {"SeqId": "2000-2", "Gene": "GF", "UniProt": "P00006", "Ensembl": ens("GF"), "Included": "Yes"},
        {"SeqId": "3000-3", "Gene": "GA, GD", "UniProt": "P00001, P00004", "Ensembl": f"{ens('GA')}, {ens('GD')}",
         "Included": "Yes"},
    ])
    decode_st02: list[dict] = field(default_factory=lambda: [
        {"gene": "GB", "SeqId": "1000-1", "cis_trans": "cis"}, {"gene": "GF", "SeqId": "2000-2", "cis_trans": "cis"},
        {"gene": "GA.GD", "SeqId": "3000-3", "cis_trans": "cis"}, {"gene": "GB", "SeqId": "1000-1", "cis_trans": "trans"},
    ])
    interval_st4: list[dict] = field(default_factory=lambda: [
        {"somamer_id": "GB.1", "UniProt": "P00002", "cis_trans": "cis"},
        {"somamer_id": "GE.1", "UniProt": "P00005 P00009", "cis_trans": "cis"},
    ])
    epigraphdb: tuple[str, ...] = ("GE",)
    frozen_candidates: list[dict] = field(default_factory=lambda: [
        {"disease": "Pilot Disease", "ot_disease_id": "pilot disease"}])
    classification_v5: list[dict] = field(default_factory=lambda: [{"gene": "GA", "disease": "Pilot Disease"}])
    classification_v5_1: list[dict] = field(default_factory=lambda: [{"gene": "GB", "disease": "pilot-disease"}])
    karim_st17: list[dict] = field(default_factory=lambda: [{"gene": "GA", "indication_mesh_id": "D000001"}])
    publication_dates: dict = field(default_factory=lambda: dict(DATES))


def hgnc_frame() -> pd.DataFrame:
    return pd.DataFrame([{"symbol": s, "status": "Approved", "ensembl_gene_id": e, "uniprot_ids": u}
                         for s, (e, u) in GENES.items()])


def _gcat_frames(w: World) -> list[tuple[pd.DataFrame, pd.DataFrame, str]]:
    st = pd.DataFrame([s for s, _ in w.gwas_catalog])
    an = pd.DataFrame([a for _, a in w.gwas_catalog])
    ust = pd.DataFrame([{"STUDY ACCESSION": "GCSTU1", "TRAIT": "unmapped", "MAPPED_TRAIT_URI": "",
                         "SUMMARY STATS LOCATION": ""}])
    uan = pd.DataFrame([{"STUDY ACCESSION": "GCSTU1", "STAGE": "initial", "BROAD ANCESTRAL CATEGORY": "European",
                         "COUNTRY OF RECRUITMENT": "", "NUMBER OF CASES": "", "NUMBER OF CONTROLS": ""}])
    return [(st, an, "published"), (ust, uan, "unpublished")]


def world_inputs(w: World) -> StageAInputs:
    disease = pd.DataFrame(w.diseases, columns=DISEASE_COLUMNS)
    hg = hgnc_frame()
    ens2sym = dict(zip(hg["ensembl_gene_id"], hg["symbol"]))
    uni2sym = {u: {s} for s, u in zip(hg["symbol"], hg["uniprot_ids"])}
    ci = pd.DataFrame(w.clinical, columns=["drugId", "diseaseId", "stage"])
    ci = pd.DataFrame({"drugId": ci["drugId"], "diseaseId": ci["diseaseId"], "phase2plus": ci["stage"].isin(PHASE2_PLUS)})
    mech = pd.DataFrame(w.mechanisms).explode("chemblIds").explode("targets").rename(
        columns={"chemblIds": "moa_drug", "targets": "ensembl"}).reset_index(drop=True)
    ukb_s, ukb_i = ukbppp_lists(list(w.ukbppp_cis), pd.DataFrame(w.olink))
    dec_s, dec_i = decode_lists(pd.DataFrame(w.decode_st01), pd.DataFrame(w.decode_st02))
    itv_s, itv_i = interval_lists(pd.DataFrame(w.interval_st4), uni2sym, set(w.epigraphdb))
    gcat = gwas_catalog_candidates(_gcat_frames(w))
    fgn = finngen_candidates(pd.DataFrame(w.finngen))
    ogw = opengwas_candidates(w.opengwas, gcat, fgn, label_index(disease))
    label_map = pilot_label_map(pd.DataFrame(w.frozen_candidates), disease)
    return StageAInputs(
        clinical_indication=ci, mechanisms=mech, molecules=pd.DataFrame(w.molecules), disease=disease,
        ens2sym=ens2sym, hpa_blood_ensembl={ens(g) for g in w.hpa_blood},
        uniprot_secreted_symbols=set(w.uniprot_secreted),
        strict_lists=InstrumentLists(ukbppp=ukb_s, decode=dec_s, interval=itv_s),
        inclusive_lists=InstrumentLists(ukbppp=ukb_i, decode=dec_i, interval=itv_i),
        candidates=combine_candidates(gcat, fgn, ogw),
        pilot_keys=pilot_keys([pd.DataFrame(w.classification_v5), pd.DataFrame(w.classification_v5_1)], label_map),
        pilot_indications=set(label_map.values()),
        karim_keys=karim_launched_keys(pd.DataFrame(w.karim_st17), disease),
        publication_dates=w.publication_dates,
        decode_diagnostic=decode_strict_diagnostic(pd.DataFrame(w.decode_st01), pd.DataFrame(w.decode_st02)),
    )


# ---- serialization to the real file formats --------------------------------------------

def _xlsx(path: Path, sheets: dict[str, tuple[int, list[str], list[list]]]) -> None:
    """sheets: name -> (header row, header, rows); rows above the header hold a title line."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, (header_row, header, rows) in sheets.items():
        ws = wb.create_sheet(name)
        ws.cell(row=1, column=1, value=f"synthetic {name}")
        for j, h in enumerate(header, start=1):
            ws.cell(row=header_row, column=j, value=h)
        for i, r in enumerate(rows, start=header_row + 1):
            for j, v in enumerate(r, start=1):
                ws.cell(row=i, column=j, value=v)
    wb.save(path)


def write_world(w: World, root: Path) -> Path:
    """Write every input file under `root`; return a paths JSON for `--paths-json`."""
    root.mkdir(parents=True, exist_ok=True)
    ot = root / "ot"
    ot.mkdir()
    ci = pd.DataFrame(w.clinical, columns=["drugId", "diseaseId", "maxClinicalStage"])
    ci["id"] = [f"CI{i}" for i in range(len(ci))]
    ci["clinicalReportIds"] = [["R1"]] * len(ci)
    pq.write_table(pa.Table.from_pandas(ci, preserve_index=False), ot / "clinical_indication.parquet")
    mech = pd.DataFrame(w.mechanisms)
    mech["targetName"] = "t"
    pq.write_table(pa.Table.from_pandas(mech, preserve_index=False), ot / "drug_mechanism_of_action.parquet")
    mol = pd.DataFrame(w.molecules)
    mol["maximumClinicalStage"] = "APPROVAL"
    pq.write_table(pa.Table.from_pandas(mol, preserve_index=False), ot / "drug_molecule.parquet")
    pq.write_table(pa.Table.from_pylist(w.diseases), ot / "disease.parquet")
    pq.write_table(pa.Table.from_pylist(w.finngen), ot / "study.parquet")

    hgnc_frame().to_csv(root / "hgnc.txt", sep="\t", index=False)
    hpa = pd.DataFrame([{"Gene": s, "Ensembl": e, "Secretome location": "Secreted to blood" if s in w.hpa_blood
                         else "Intracellular and membrane"} for s, (e, _) in GENES.items()])
    buf = io.StringIO()
    hpa.to_csv(buf, sep="\t", index=False)
    with zipfile.ZipFile(root / "proteinatlas.tsv.zip", "w") as z:
        z.writestr("proteinatlas.tsv", buf.getvalue())
    for tag in ("KW-0964", "SL-0243"):
        pd.DataFrame([{"Entry": GENES[g][1], "Gene Names (primary)": g} for g in w.uniprot_secreted]).to_csv(
            root / f"uniprot_{tag}.tsv", sep="\t", index=False)
    gc_paths = {}
    for (st, an, kind), (sk, ak) in zip(_gcat_frames(w), (("studies", "ancestries"),
                                                          ("unpublished_studies", "unpublished_ancestries"))):
        st.to_csv(root / f"gc_{sk}.tsv", sep="\t", index=False)
        an.to_csv(root / f"gc_{ak}.tsv", sep="\t", index=False)
        gc_paths[sk], gc_paths[ak] = root / f"gc_{sk}.tsv", root / f"gc_{ak}.tsv"
    (root / "gwasinfo.json").write_text(json.dumps(w.opengwas))
    (root / "ukbppp_cis.json").write_text(json.dumps(list(w.ukbppp_cis)))
    pd.DataFrame(w.olink).to_csv(root / "olink.tsv", sep="\t", index=False)
    _xlsx(root / "decode.xlsx", {
        "ST01": (3, ["SeqId", "Protein (short name)", "Gene", "Included in\nanalysis", "UniProt", "Ensembl.Gene.ID"],
                 [[r["SeqId"], "p", r["Gene"], r["Included"], r["UniProt"], r["Ensembl"]] for r in w.decode_st01]),
        "ST02": (3, ["pQTL_ID \n(global)", "gene\n (prot.)", "UniProt", "SeqId", "cis/\ntrans"],
                 [[i, r["gene"], "", r["SeqId"], r["cis_trans"]] for i, r in enumerate(w.decode_st02)]),
    })
    _xlsx(root / "sun2018.xlsx", {"ST4 - pQTL summary": (
        5, ["Locus ID", "SOMAmer ID", "Target", "UniProt", "cis/ trans"],
        [[i, r["somamer_id"], "t", r["UniProt"], r["cis_trans"]] for i, r in enumerate(w.interval_st4)])})
    pd.DataFrame({"gene_symbol": [";".join(w.epigraphdb)], "mr_pval": ["not read"]}).to_csv(
        root / "epigraphdb.csv", index=False)
    pd.DataFrame(w.frozen_candidates).to_csv(root / "frozen.csv", index=False)
    pd.DataFrame(w.classification_v5).assign(outcome="not read").to_csv(root / "class_v5.csv", index=False)
    pd.DataFrame(w.classification_v5_1).assign(outcome="not read").to_csv(root / "class_v5_1.csv", index=False)
    _xlsx(root / "karim.xlsx", {"ST17 - pqtl_success_ti_pairs": (
        3, ["gene", "indication_mesh_id", "indication_mesh_term", "therapeutic_area", "ti_uid"],
        [[r["gene"], r["indication_mesh_id"], "term", "ta", "uid"] for r in w.karim_st17])})

    paths = {
        "ot_clinical_indication": ot / "clinical_indication.parquet",
        "ot_drug_mechanism_of_action": ot / "drug_mechanism_of_action.parquet",
        "ot_drug_molecule": ot / "drug_molecule.parquet", "ot_disease": ot / "disease.parquet",
        "ot_study": ot / "study.parquet", "hgnc": root / "hgnc.txt", "hpa": root / "proteinatlas.tsv.zip",
        "uniprot_kw0964": root / "uniprot_KW-0964.tsv", "uniprot_sl0243": root / "uniprot_SL-0243.tsv",
        "gwas_catalog_studies": gc_paths["studies"], "gwas_catalog_ancestries": gc_paths["ancestries"],
        "gwas_catalog_unpublished_studies": gc_paths["unpublished_studies"],
        "gwas_catalog_unpublished_ancestries": gc_paths["unpublished_ancestries"],
        "opengwas_gwasinfo": root / "gwasinfo.json", "ukbppp_cis_list": root / "ukbppp_cis.json",
        "olink_map": root / "olink.tsv", "decode_supplement": root / "decode.xlsx",
        "interval_supplement": root / "sun2018.xlsx", "epigraphdb_catalog": root / "epigraphdb.csv",
        "frozen_candidates_v34": root / "frozen.csv", "classification_v5": root / "class_v5.csv",
        "classification_v5_1": root / "class_v5_1.csv", "karim2026_workbook": root / "karim.xlsx",
    }
    out = root / "paths.json"
    out.write_text(json.dumps({k: str(v) for k, v in paths.items()}))
    return out
